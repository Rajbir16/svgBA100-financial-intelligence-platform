"""Sprint 4 Valuation Module for Nifty 100.

Computes FCF Yield, 5-year median P/E, sector-relative P/E benchmarking,
and assigns valuation flags (Caution, Discount, Fair).
Generates output/valuation_summary.xlsx and output/valuation_flags.csv.
"""

import os
import sqlite3
from typing import Any, Dict, Optional
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def compute_valuation_data(db_path: str = "nifty100.db") -> pd.DataFrame:
    """Compute valuation metrics and flags for all 92 companies."""
    conn = sqlite3.connect(db_path)
    
    # Query latest annual FCF from financial_ratios and 2024 market cap metrics
    query = """
        WITH latest_fr AS (
            SELECT company_id, free_cash_flow_cr,
                   ROW_NUMBER() OVER (
                       PARTITION BY company_id
                       ORDER BY 
                           CASE 
                               WHEN company_id = 'SIEMENS' AND year = 'Sep 2024' THEN 1
                               WHEN company_id != 'SIEMENS' AND year = 'Mar 2024' THEN 1
                               WHEN year LIKE '%2024%' AND year != 'TTM' THEN 2
                               ELSE 3
                           END,
                           id DESC
                   ) as rn
            FROM financial_ratios
            WHERE sales_cr IS NOT NULL AND year != 'TTM'
        )
        SELECT 
            c.id AS company_id,
            c.company_name,
            s.broad_sector AS sector,
            mc.pe_ratio AS pe,
            mc.pb_ratio AS pb,
            mc.ev_ebitda,
            mc.market_cap_crore AS market_cap,
            fr.free_cash_flow_cr AS fcf
        FROM companies c
        LEFT JOIN sectors s ON c.id = s.company_id
        LEFT JOIN market_cap mc ON c.id = mc.company_id AND mc.year = 2024
        LEFT JOIN latest_fr fr ON c.id = fr.company_id AND fr.rn = 1
        ORDER BY c.id
    """
    df = pd.read_sql_query(query, conn)

    # 5-year historical median PE (2020 to 2024)
    hist_pe_df = pd.read_sql_query(
        "SELECT company_id, pe_ratio FROM market_cap WHERE year BETWEEN 2020 AND 2024 AND pe_ratio IS NOT NULL",
        conn
    )
    conn.close()

    # Calculate 5-year median PE per company
    hist_med = hist_pe_df.groupby("company_id")["pe_ratio"].median().to_dict()
    df["5yr_median_PE"] = df["company_id"].map(hist_med).round(2)

    # Calculate sector median PE using 2024 values
    valid_pe = df[df["pe"].notna() & (df["pe"] > 0)]
    sector_medians = valid_pe.groupby("sector")["pe"].median().to_dict()
    df["sector_median_PE"] = df["sector"].map(sector_medians)

    # FCF Yield (%) = FCF / Market Cap Crore * 100
    def calc_fcf_yield(row: pd.Series) -> Optional[float]:
        fcf = row["fcf"]
        mcap = row["market_cap"]
        if pd.isna(fcf) or pd.isna(mcap) or mcap is None or mcap <= 0:
            return None
        return round((float(fcf) / float(mcap)) * 100.0, 2)

    df["FCF_yield_pct"] = df.apply(calc_fcf_yield, axis=1)

    # PE vs Sector Median (%) = ((P/E - sector_median) / sector_median) * 100
    def calc_pe_vs_sector(row: pd.Series) -> Optional[float]:
        pe = row["pe"]
        sec_med = row["sector_median_PE"]
        if pd.isna(pe) or pd.isna(sec_med) or sec_med is None or sec_med <= 0:
            return None
        return round(((float(pe) - float(sec_med)) / float(sec_med)) * 100.0, 2)

    df["PE_vs_sector_median_pct"] = df.apply(calc_pe_vs_sector, axis=1)

    # Valuation Flag logic:
    # If P/E > sector_median * 1.5 -> Caution
    # If P/E < sector_median * 0.7 -> Discount
    # Else -> Fair
    def assign_flag(row: pd.Series) -> str:
        pe = row["pe"]
        sec_med = row["sector_median_PE"]
        if pd.isna(pe) or pd.isna(sec_med) or sec_med is None or sec_med <= 0:
            return "Fair"
        if float(pe) > float(sec_med) * 1.5:
            return "Caution"
        elif float(pe) < float(sec_med) * 0.7:
            return "Discount"
        else:
            return "Fair"

    df["flag"] = df.apply(assign_flag, axis=1)

    # Clean up column names to match required specifications exactly:
    # company_id, company_name, sector, P/E, P/B, EV/EBITDA, FCF_yield_pct, 5yr_median_PE, PE_vs_sector_median_pct, flag
    df["P/E"] = df["pe"].round(2)
    df["P/B"] = df["pb"].round(2)
    df["EV/EBITDA"] = df["ev_ebitda"].round(2)

    return df


def generate_valuation_reports(db_path: str = "nifty100.db", output_dir: str = "output") -> Dict[str, Any]:
    """Generate output/valuation_summary.xlsx and output/valuation_flags.csv."""
    os.makedirs(output_dir, exist_ok=True)
    df = compute_valuation_data(db_path)

    summary_cols = [
        "company_id",
        "company_name",
        "sector",
        "P/E",
        "P/B",
        "EV/EBITDA",
        "FCF_yield_pct",
        "5yr_median_PE",
        "PE_vs_sector_median_pct",
        "flag"
    ]
    summary_df = df[summary_cols].copy()

    # 1. Generate output/valuation_summary.xlsx (92 rows)
    xlsx_path = os.path.join(output_dir, "valuation_summary.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valuation Summary"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    caution_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    caution_font = Font(name="Calibri", size=10, bold=True, color="721C24")
    discount_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    discount_font = Font(name="Calibri", size=10, bold=True, color="155724")
    fair_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    fair_font = Font(name="Calibri", size=10, color="495057")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )

    ws.append(summary_cols)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(summary_cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for r_idx, (_, row) in enumerate(summary_df.iterrows(), start=2):
        row_vals = []
        for col in summary_cols:
            val = row[col]
            if pd.isna(val) or val is None:
                row_vals.append("N/A")
            elif isinstance(val, (float, np.floating)):
                row_vals.append(round(float(val), 2))
            else:
                row_vals.append(str(val))

        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 20

        flag_val = str(row["flag"])
        for c_idx in range(1, len(row_vals) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.font = data_font
            cell.alignment = Alignment(horizontal="center" if c_idx not in (2, 3) else "left", vertical="center")

            # Flag cell styling
            if c_idx == len(summary_cols):
                if flag_val == "Caution":
                    cell.fill = caution_fill
                    cell.font = caution_font
                elif flag_val == "Discount":
                    cell.fill = discount_fill
                    cell.font = discount_font
                else:
                    cell.fill = fair_fill
                    cell.font = fair_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(xlsx_path)

    # 2. Generate output/valuation_flags.csv containing only Caution and Discount companies
    csv_path = os.path.join(output_dir, "valuation_flags.csv")
    flags_df = summary_df[summary_df["flag"].isin(["Caution", "Discount"])].copy()
    flags_df.to_csv(csv_path, index=False, encoding="utf-8")

    return {
        "summary_rows": len(summary_df),
        "flags_rows": len(flags_df),
        "caution_count": int((flags_df["flag"] == "Caution").sum()),
        "discount_count": int((flags_df["flag"] == "Discount").sum()),
        "xlsx_path": xlsx_path,
        "csv_path": csv_path
    }


def get_company_valuation(ticker: str, db_path: str = "nifty100.db") -> Optional[Dict[str, Any]]:
    """Return valuation dictionary for a specific company ticker."""
    df = compute_valuation_data(db_path)
    match = df[df["company_id"].str.upper() == ticker.strip().upper()]
    if match.empty:
        return None
    return match.iloc[0].to_dict()


if __name__ == "__main__":
    res = generate_valuation_reports()
    print("Valuation reports generated successfully:")
    print(res)

