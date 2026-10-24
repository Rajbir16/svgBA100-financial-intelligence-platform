"""Sprint 5 Day 29 — NLP Analysis Text Parser.

Parses unstructured multi-period financial text blocks from analysis records
using regular expressions, normalizes periods and percentage values, records
failures with actionable diagnostic reasons, and cross-validates parsed CAGRs
against the Ratio Engine calculations.
"""

import argparse
import glob
import os
import re
import sqlite3
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

# Regular expression compatible with text such as '5 Years: 18.5%'
# Handles:
# - Spacing variations before/after digits and colons
# - Singular and plural period labels ('Year' vs 'Years')
# - Optional punctuation (colons ':', dashes '-', whitespace gaps)
# - Negative percentage values ('-2%') without swallowing the minus sign
# - Integer and floating point decimal percentage values ('18.5%', '21%')
ANALYSIS_YEARS_REGEX = re.compile(
    r"^\s*(\d+)\s*years?(?:[:\s\t]+|(?:\s*-\s*))\s*(-?\d+(?:\.\d+)?)\s*%\s*$",
    re.IGNORECASE,
)

ANALYSIS_FIELDS = [
    "compounded_sales_growth",
    "compounded_profit_growth",
    "stock_price_cagr",
    "roe",
]


def parse_analysis_text(text: Any) -> Optional[Tuple[int, float]]:
    """Parse text formatted like '5 Years: 18.5%' into (period_years, value_pct).

    Args:
        text: Raw text string (e.g. '10 Years: 21%', '5 Years 14%', '1 Year: -2%').

    Returns:
        Tuple of (period_years: int, value_pct: float) if matched, else None.
    """
    if text is None or pd.isna(text):
        return None
    cleaned = str(text).strip()
    if not cleaned:
        return None
    match = ANALYSIS_YEARS_REGEX.match(cleaned)
    if match:
        period_years = int(match.group(1))
        value_pct = float(match.group(2))
        return period_years, value_pct
    return None


def parse_analysis_data(
    db_path: str = "nifty100.db",
    excel_path: Optional[str] = None,
    output_dir: str = "output",
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """Parse analysis records from SQLite or Excel file.

    Validates company IDs against the master companies table.
    Outputs parsed records to 'output/analysis_parsed.csv' and failures to
    'output/parse_failures.csv'.

    Args:
        db_path: Path to nifty100.db SQLite database.
        excel_path: Optional path to raw analysis.xlsx file.
        output_dir: Destination directory for output CSVs.

    Returns:
        Tuple of (parsed_df, failures_df).
    """
    os.makedirs(output_dir, exist_ok=True)

    # 1. Fetch valid company IDs from SQLite master table
    valid_companies: set = set()
    if os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        try:
            cur = conn.cursor()
            cur.execute("SELECT id FROM companies")
            valid_companies = {row[0] for row in cur.fetchall()}
        finally:
            conn.close()

    # 2. Load analysis records
    rows_to_process: List[Dict[str, Any]] = []

    # If excel_path is provided or discoverable in Dataset directory, inspect raw Excel
    candidate_excel = excel_path
    if not candidate_excel:
        default_dir = "./data"
        if os.path.exists(default_dir):
            matches = glob.glob(os.path.join(default_dir, "*analysis*.xlsx"))
            if matches:
                candidate_excel = matches[0]

    if candidate_excel and os.path.exists(candidate_excel):
        import openpyxl

        wb = openpyxl.load_workbook(candidate_excel, data_only=True)
        sheet = wb["Analysis"] if "Analysis" in wb.sheetnames else wb.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        # Header is typically at row 1 or 2
        header_idx = 0
        for idx, r in enumerate(raw_rows[:5]):
            if r and any(
                isinstance(cell, str) and "company_id" in cell.lower() for cell in r
            ):
                header_idx = idx
                break
        headers = [str(c).strip().lower() if c else "" for c in raw_rows[header_idx]]
        for r in raw_rows[header_idx + 1 :]:
            if not r or not any(r):
                continue
            r_dict = dict(zip(headers, r))
            rows_to_process.append(r_dict)
    elif os.path.exists(db_path):
        conn = sqlite3.connect(db_path)
        try:
            df_db = pd.read_sql("SELECT * FROM analysis", conn)
            rows_to_process = df_db.to_dict(orient="records")
        finally:
            conn.close()

    parsed_records: List[Dict[str, Any]] = []
    failure_records: List[Dict[str, Any]] = []

    for row in rows_to_process:
        raw_cid = row.get("company_id")
        cid = str(raw_cid).strip().upper() if raw_cid is not None else ""

        # Validate company_id against master companies
        if valid_companies and cid not in valid_companies:
            for field in ANALYSIS_FIELDS:
                val = row.get(field)
                if val is not None and str(val).strip():
                    failure_records.append(
                        {
                            "company_id": cid or "UNKNOWN",
                            "source_field": field,
                            "original_text": str(val).strip(),
                            "failure_reason": (
                                f"Unmapped company ID '{cid}': not found in master companies table"
                            ),
                        }
                    )
            continue

        # Parse each of the 4 target fields
        for field in ANALYSIS_FIELDS:
            raw_text = row.get(field)
            if raw_text is None or pd.isna(raw_text):
                failure_records.append(
                    {
                        "company_id": cid,
                        "source_field": field,
                        "original_text": "NULL",
                        "failure_reason": "Missing / null field value",
                    }
                )
                continue

            cleaned_text = str(raw_text).strip()
            if not cleaned_text:
                failure_records.append(
                    {
                        "company_id": cid,
                        "source_field": field,
                        "original_text": "",
                        "failure_reason": "Empty string value",
                    }
                )
                continue

            parsed = parse_analysis_text(cleaned_text)
            if parsed is not None:
                period_years, value_pct = parsed
                parsed_records.append(
                    {
                        "company_id": cid,
                        "metric_type": field,
                        "period_years": period_years,
                        "value_pct": value_pct,
                    }
                )
            else:
                # Classify non-matching reason
                if "ttm" in cleaned_text.lower():
                    reason = "Non-standard period format: TTM (Trailing Twelve Months) lacks discrete year count"
                elif "last year" in cleaned_text.lower():
                    reason = "Non-standard period format: 'Last Year' lacks numeric year count"
                elif "%" not in cleaned_text:
                    reason = "Malformed text: missing percentage symbol '%'"
                else:
                    reason = "Unmatched text: does not conform to 'N Year(s): X%' regex pattern"

                failure_records.append(
                    {
                        "company_id": cid,
                        "source_field": field,
                        "original_text": cleaned_text,
                        "failure_reason": reason,
                    }
                )

    # Build DataFrames and deduplicate
    parsed_df = pd.DataFrame(parsed_records)
    if not parsed_df.empty:
        parsed_df = parsed_df.drop_duplicates(
            subset=["company_id", "metric_type", "period_years"]
        )
        parsed_df = parsed_df.sort_values(
            by=["company_id", "metric_type", "period_years"]
        )
    else:
        parsed_df = pd.DataFrame(
            columns=["company_id", "metric_type", "period_years", "value_pct"]
        )

    failures_df = pd.DataFrame(failure_records)
    if not failures_df.empty:
        failures_df = failures_df.drop_duplicates()
        failures_df = failures_df.sort_values(by=["company_id", "source_field"])
    else:
        failures_df = pd.DataFrame(
            columns=[
                "company_id",
                "source_field",
                "original_text",
                "failure_reason",
            ]
        )

    # Export CSVs
    parsed_path = os.path.join(output_dir, "analysis_parsed.csv")
    failures_path = os.path.join(output_dir, "parse_failures.csv")

    parsed_df.to_csv(parsed_path, index=False)
    failures_df.to_csv(failures_path, index=False)

    return parsed_df, failures_df


def cross_validate_cagrs(
    parsed_df: pd.DataFrame,
    db_path: str = "nifty100.db",
    output_dir: str = "output",
    divergence_threshold_pp: float = 5.0,
) -> pd.DataFrame:
    """Cross-validate parsed sales and profit CAGRs against Ratio Engine calculations.

    Flags absolute divergence greater than 5 percentage points for manual review.
    Compares only strictly equivalent metrics and periods (3yr, 5yr, 10yr).

    Methodology:
    - Parsed 'compounded_sales_growth' with period P is mapped to 'revenue_cagr_{P}yr'.
    - Parsed 'compounded_profit_growth' with period P is mapped to 'pat_cagr_{P}yr'.
    - Both metrics represent annualized percentage compound growth rates (CAGR %).
    - Divergence = abs(parsed_pct - calculated_pct).
    - If divergence > 5.0 percentage points, flag_divergence = True.

    Args:
        parsed_df: DataFrame containing parsed records.
        db_path: Path to nifty100.db SQLite database.
        output_dir: Directory where validation report CSV will be written.
        divergence_threshold_pp: Threshold in percentage points to flag for review.

    Returns:
        DataFrame with cross-validation comparisons.
    """
    if not os.path.exists(db_path):
        raise FileNotFoundError(f"Database not found: {db_path}")

    conn = sqlite3.connect(db_path)
    try:
        # Load latest authoritative financial ratios
        # Ranked by latest available annual year per company
        query = """
            WITH ranked_fr AS (
                SELECT fr.*,
                       ROW_NUMBER() OVER (
                           PARTITION BY fr.company_id
                           ORDER BY 
                               CASE 
                                   WHEN fr.company_id = 'SIEMENS' AND fr.year = 'Sep 2024' THEN 1
                                   WHEN fr.company_id != 'SIEMENS' AND fr.year = 'Mar 2024' THEN 1
                                   WHEN fr.year LIKE '%2024%' AND fr.year != 'TTM' THEN 2
                                   ELSE 3
                               END,
                               fr.id DESC
                       ) as rn
                FROM financial_ratios fr
                WHERE fr.year != 'TTM'
            )
            SELECT company_id, year,
                   revenue_cagr_3yr, revenue_cagr_5yr, revenue_cagr_10yr,
                   pat_cagr_3yr, pat_cagr_5yr, pat_cagr_10yr
            FROM ranked_fr
            WHERE rn = 1
        """
        fr_df = pd.read_sql(query, conn).set_index("company_id")
    finally:
        conn.close()

    comparisons: List[Dict[str, Any]] = []

    metric_mapping = {
        "compounded_sales_growth": "revenue_cagr",
        "compounded_profit_growth": "pat_cagr",
    }

    for _, row in parsed_df.iterrows():
        cid = row["company_id"]
        metric_type = row["metric_type"]
        period = row["period_years"]
        parsed_val = row["value_pct"]

        if metric_type not in metric_mapping:
            continue  # Stock price CAGR and ROE are not direct sales/PAT CAGRs

        if cid not in fr_df.index:
            continue

        base_metric = metric_mapping[metric_type]
        calc_col = f"{base_metric}_{period}yr"

        calc_row = fr_df.loc[cid]
        if calc_col not in calc_row:
            continue

        calc_val = calc_row[calc_col]
        if pd.isna(calc_val) or calc_val is None:
            continue

        calc_float = float(calc_val)
        abs_diff = round(abs(parsed_val - calc_float), 4)
        is_flagged = abs_diff > divergence_threshold_pp

        comparisons.append(
            {
                "company_id": cid,
                "parsed_metric": metric_type,
                "engine_metric": calc_col,
                "period_years": period,
                "parsed_value_pct": parsed_val,
                "engine_value_pct": round(calc_float, 4),
                "absolute_divergence_pp": abs_diff,
                "divergence_gt_5pp": is_flagged,
                "status": "FLAG_FOR_REVIEW" if is_flagged else "ALIGNED",
            }
        )

    val_df = pd.DataFrame(comparisons)
    if not val_df.empty:
        val_df = val_df.sort_values(
            by=["divergence_gt_5pp", "absolute_divergence_pp"], ascending=[False, False]
        )
        report_path = os.path.join(output_dir, "cagr_cross_validation.csv")
        val_df.to_csv(report_path, index=False)

    return val_df


def main():
    parser = argparse.ArgumentParser(description="Sprint 5 NLP Analysis Text Parser")
    parser.add_argument("--db", default="nifty100.db", help="Path to SQLite database")
    parser.add_argument("--excel", default=None, help="Path to source analysis.xlsx")
    parser.add_argument("--output-dir", default="output", help="Output directory")

    args = parser.parse_args()

    print("Running NLP Analysis Text Parser...")
    parsed_df, failures_df = parse_analysis_data(
        db_path=args.db, excel_path=args.excel, output_dir=args.output_dir
    )
    print(f"Parsed records generated: {len(parsed_df)}")
    print(f"Parse failures recorded: {len(failures_df)}")

    print("\nRunning Ratio Engine CAGR Cross-Validation...")
    val_df = cross_validate_cagrs(
        parsed_df, db_path=args.db, output_dir=args.output_dir
    )
    flagged = val_df[val_df["divergence_gt_5pp"] == True]
    print(f"Total comparable CAGR pairs: {len(val_df)}")
    print(f"Pairs flagged with divergence > 5 pp: {len(flagged)}")
    if not flagged.empty:
        print("\nFlagged entries for manual review:")
        for _, r in flagged.iterrows():
            print(
                f"  {r['company_id']} | {r['parsed_metric']} ({r['period_years']}Y) | "
                f"Parsed={r['parsed_value_pct']}% vs Engine={r['engine_value_pct']}% | "
                f"Diff={r['absolute_divergence_pp']} pp"
            )


if __name__ == "__main__":
    main()

