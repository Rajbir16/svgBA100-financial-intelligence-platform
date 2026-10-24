"""Excel ETL Loader for svgBA100 Financial Intelligence Platform.

Reads 12 financial datasets from Excel, normalizes fields, enforces
data quality rules, loads records into SQLite with full transaction support,
rejects referential violations, and produces audit reports.
"""

import argparse
import csv
import fnmatch
import os
import sqlite3
from datetime import datetime
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl

from src.etl.normaliser import (
    normalize_boolean,
    normalize_date,
    normalize_numeric,
    normalize_string,
    normalize_ticker,
    normalize_year,
)
from src.etl.validator import DataValidator

TABLE_CONFIGS = [
    {
        "table_name": "companies",
        "file_pattern": "*companies.xlsx",
        "sheet_name": "Companies",
        "is_parent": True,
        "columns": [
            ("id", "id", "ticker"),
            ("company_logo", "company_logo", "string"),
            ("company_name", "company_name", "string"),
            ("chart_link", "chart_link", "string"),
            ("about_company", "about_company", "string"),
            ("website", "website", "string"),
            ("nse_profile", "nse_profile", "string"),
            ("bse_profile", "bse_profile", "string"),
            ("face_value", "face_value", "numeric"),
            ("book_value", "book_value", "numeric"),
            ("roce_percentage", "roce_percentage", "numeric"),
            ("roe_percentage", "roe_percentage", "numeric"),
        ],
    },
    {
        "table_name": "sectors",
        "file_pattern": "*sectors.xlsx",
        "sheet_name": "Sheet1",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("broad_sector", "broad_sector", "string"),
            ("sub_sector", "sub_sector", "string"),
            ("index_weight_pct", "index_weight_pct", "numeric"),
            ("market_cap_category", "market_cap_category", "string"),
        ],
    },
    {
        "table_name": "market_cap",
        "file_pattern": "*market_cap.xlsx",
        "sheet_name": "Sheet1",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("year", "year", "int_year"),
            ("market_cap_crore", "market_cap_crore", "numeric"),
            ("enterprise_value_crore", "enterprise_value_crore", "numeric"),
            ("pe_ratio", "pe_ratio", "numeric"),
            ("pb_ratio", "pb_ratio", "numeric"),
            ("ev_ebitda", "ev_ebitda", "numeric"),
            ("dividend_yield_pct", "dividend_yield_pct", "numeric"),
        ],
    },
    {
        "table_name": "stock_prices",
        "file_pattern": "*stock_prices.xlsx",
        "sheet_name": "Sheet1",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("date", "date", "date"),
            ("open_price", "open_price", "numeric"),
            ("high_price", "high_price", "numeric"),
            ("low_price", "low_price", "numeric"),
            ("close_price", "close_price", "numeric"),
            ("volume", "volume", "int_numeric"),
            ("adjusted_close", "adjusted_close", "numeric"),
        ],
    },
    {
        "table_name": "financial_ratios",
        "file_pattern": "*financial_ratios.xlsx",
        "sheet_name": "Sheet1",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("year", "year", "year"),
            ("net_profit_margin_pct", "net_profit_margin_pct", "numeric"),
            ("operating_profit_margin_pct", "operating_profit_margin_pct", "numeric"),
            ("return_on_equity_pct", "return_on_equity_pct", "numeric"),
            ("debt_to_equity", "debt_to_equity", "numeric"),
            ("interest_coverage", "interest_coverage", "numeric"),
            ("asset_turnover", "asset_turnover", "numeric"),
            ("free_cash_flow_cr", "free_cash_flow_cr", "numeric"),
            ("capex_cr", "capex_cr", "numeric"),
            ("earnings_per_share", "earnings_per_share", "numeric"),
            ("book_value_per_share", "book_value_per_share", "numeric"),
            ("dividend_payout_ratio_pct", "dividend_payout_ratio_pct", "numeric"),
            ("total_debt_cr", "total_debt_cr", "numeric"),
            ("cash_from_operations_cr", "cash_from_operations_cr", "numeric"),
        ],
    },
    {
        "table_name": "balancesheet",
        "file_pattern": "*balancesheet.xlsx",
        "sheet_name": "Balance Sheet",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("year", "year", "year"),
            ("equity_capital", "equity_capital", "numeric"),
            ("reserves", "reserves", "numeric"),
            ("borrowings", "borrowings", "numeric"),
            ("other_liabilities", "other_liabilities", "numeric"),
            ("total_liabilities", "total_liabilities", "numeric"),
            ("fixed_assets", "fixed_assets", "numeric"),
            ("cwip", "cwip", "numeric"),
            ("investments", "investments", "numeric"),
            ("other_asset", "other_asset", "numeric"),
            ("total_assets", "total_assets", "numeric"),
        ],
    },
    {
        "table_name": "profitandloss",
        "file_pattern": "*profitandloss.xlsx",
        "sheet_name": "Profit & Loss",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("year", "year", "year"),
            ("sales", "sales", "numeric"),
            ("expenses", "expenses", "numeric"),
            ("operating_profit", "operating_profit", "numeric"),
            ("opm_percentage", "opm_percentage", "numeric"),
            ("other_income", "other_income", "numeric"),
            ("interest", "interest", "numeric"),
            ("depreciation", "depreciation", "numeric"),
            ("profit_before_tax", "profit_before_tax", "numeric"),
            ("tax_percentage", "tax_percentage", "numeric"),
            ("net_profit", "net_profit", "numeric"),
            ("eps", "eps", "numeric"),
            ("dividend_payout", "dividend_payout", "numeric"),
        ],
    },
    {
        "table_name": "cashflow",
        "file_pattern": "*cashflow.xlsx",
        "sheet_name": "Cash Flow",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("year", "year", "year"),
            ("operating_activity", "operating_activity", "numeric"),
            ("investing_activity", "investing_activity", "numeric"),
            ("financing_activity", "financing_activity", "numeric"),
            ("net_cash_flow", "net_cash_flow", "numeric"),
        ],
    },
    {
        "table_name": "peer_groups",
        "file_pattern": "*peer_groups.xlsx",
        "sheet_name": "Sheet1",
        "columns": [
            ("id", "id", "int_pk"),
            ("peer_group_name", "peer_group_name", "string"),
            ("company_id", "company_id", "ticker"),
            ("is_benchmark", "is_benchmark", "boolean"),
        ],
    },
    {
        "table_name": "analysis",
        "file_pattern": "*analysis.xlsx",
        "sheet_name": "Analysis",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("compounded_sales_growth", "compounded_sales_growth", "string"),
            ("compounded_profit_growth", "compounded_profit_growth", "string"),
            ("stock_price_cagr", "stock_price_cagr", "string"),
            ("roe", "roe", "string"),
        ],
    },
    {
        "table_name": "prosandcons",
        "file_pattern": "*prosandcons.xlsx",
        "sheet_name": "Pros & Cons",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("pros", "pros", "string"),
            ("cons", "cons", "string"),
        ],
    },
    {
        "table_name": "documents",
        "file_pattern": "*documents.xlsx",
        "sheet_name": "Documents",
        "columns": [
            ("id", "id", "int_pk"),
            ("company_id", "company_id", "ticker"),
            ("Year", "year", "year"),
            ("Annual_Report", "annual_report", "string"),
        ],
    },
]


class ETLLoader:
    def __init__(
        self,
        db_path: str = "nifty100.db",
        data_dir: str = "./data",
        output_dir: str = "output",
        schema_path: str = "db/schema.sql",
    ):
        self.db_path = db_path
        self.data_dir = data_dir
        self.output_dir = output_dir
        self.schema_path = schema_path
        self.validator = DataValidator()
        self.audit_log: List[Dict[str, Any]] = []
        self.valid_company_ids: Set[str] = set()

    def find_file(self, pattern: str) -> Optional[str]:
        """Locate file matching pattern in data_dir."""
        if not os.path.exists(self.data_dir):
            raise FileNotFoundError(f"Data directory not found: {self.data_dir}")
        for fname in os.listdir(self.data_dir):
            if fnmatch.fnmatch(fname, pattern) and fname.endswith(".xlsx"):
                return os.path.join(self.data_dir, fname)
        return None

    def initialize_database(self, conn: sqlite3.Connection):
        """Execute schema.sql to initialize all tables and indexes."""
        conn.execute("PRAGMA foreign_keys = ON;")
        if os.path.exists(self.schema_path):
            with open(self.schema_path, "r", encoding="utf-8") as f:
                schema_sql = f.read()
            conn.executescript(schema_sql)
        else:
            raise FileNotFoundError(f"Schema file not found: {self.schema_path}")

    def read_sheet_data(
        self, file_path: str, sheet_name: str
    ) -> Tuple[List[str], List[Tuple[Any, ...]]]:
        """Read worksheet and detect header row (row 1 vs row 2 banner)."""
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
        else:
            sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], []

        # Check if row 1 is a title/banner row (e.g. col 1 has text, col 2 is None, and row 2 has 'id')
        if (
            len(rows) > 1
            and rows[0][1] is None
            and len(rows[1]) > 0
            and str(rows[1][0]).strip().lower() == "id"
        ):
            headers = [
                str(c).strip() if c is not None else "" for c in rows[1]
            ]
            data_rows = rows[2:]
        else:
            headers = [
                str(c).strip() if c is not None else "" for c in rows[0]
            ]
            data_rows = rows[1:]

        # Filter out completely empty rows
        filtered_data = [
            r for r in data_rows if any(c is not None and str(c).strip() != "" for c in r)
        ]
        return headers, filtered_data

    def normalize_cell(self, val: Any, col_type: str) -> Any:
        """Apply type-specific normalization."""
        if col_type == "ticker":
            return normalize_ticker(val)
        elif col_type == "int_pk":
            return normalize_numeric(val, allow_float=False)
        elif col_type == "int_year":
            num = normalize_numeric(val, allow_float=False)
            return int(num) if num is not None else None
        elif col_type == "year":
            return normalize_year(val)
        elif col_type == "date":
            return normalize_date(val)
        elif col_type == "numeric":
            return normalize_numeric(val, allow_float=True)
        elif col_type == "int_numeric":
            return normalize_numeric(val, allow_float=False)
        elif col_type == "boolean":
            return normalize_boolean(val)
        elif col_type == "string":
            return normalize_string(val)
        return normalize_string(val)

    def load_table(
        self, conn: sqlite3.Connection, config: Dict[str, Any]
    ) -> Dict[str, Any]:
        """Normalize, validate, and load a single table."""
        table_name = config["table_name"]
        file_pattern = config["file_pattern"]
        sheet_name = config["sheet_name"]
        cols_cfg = config["columns"]
        is_parent = config.get("is_parent", False)

        file_path = self.find_file(file_pattern)
        if not file_path:
            raise FileNotFoundError(
                f"Source file matching {file_pattern} not found in {self.data_dir}"
            )

        src_filename = os.path.basename(file_path)
        headers, raw_rows = self.read_sheet_data(file_path, sheet_name)

        # Build column index mapping (case-insensitive)
        header_map = {h.lower(): idx for idx, h in enumerate(headers) if h}

        target_cols = [c[1] for c in cols_cfg]
        placeholders = ", ".join(["?"] * len(target_cols))
        insert_sql = f"INSERT INTO {table_name} ({', '.join(target_cols)}) VALUES ({placeholders})"

        rows_read = len(raw_rows)
        rows_inserted = 0
        rows_rejected = 0
        seen_pks: Set[Any] = set()
        seen_pairs: Set[Tuple[str, str]] = set()

        records_to_insert = []
        pnl_records: Dict[Tuple[str, str], float] = {}

        for row_idx, r in enumerate(raw_rows, start=1):
            record: Dict[str, Any] = {}
            raw_record: Dict[str, Any] = {}

            for src_col, tgt_col, col_type in cols_cfg:
                src_idx = header_map.get(src_col.lower())
                raw_val = r[src_idx] if (src_idx is not None and src_idx < len(r)) else None
                raw_record[tgt_col] = raw_val
                record[tgt_col] = self.normalize_cell(raw_val, col_type)

                # DQ-16 check for raw placeholders
                self.validator.check_malformed_values(
                    table_name,
                    record.get("company_id") or record.get("id"),
                    record.get("year"),
                    tgt_col,
                    raw_val,
                )

                # DQ-07 check for non-numeric in numeric column
                if col_type in ("numeric", "int_numeric", "int_pk", "int_year"):
                    self.validator.check_numeric_validity(
                        table_name,
                        record.get("company_id") or record.get("id"),
                        record.get("year"),
                        tgt_col,
                        raw_val,
                        record[tgt_col],
                    )

            # --- Pre-insertion Validation ---
            pk_val = record.get("id")
            # DQ-01 & DQ-02
            is_pk_valid = self.validator.check_pk_validity(
                table_name, pk_val, is_numeric=(col_type != "ticker" and not is_parent)
            )
            is_pk_unique = self.validator.check_pk_uniqueness(
                table_name, pk_val, seen_pks
            )

            if not (is_pk_valid and is_pk_unique):
                rows_rejected += 1
                continue

            if is_parent:
                # Parent table: register valid company ID
                cid = record.get("id")
                if cid:
                    self.valid_company_ids.add(cid)
            else:
                # Child table: DQ-04 Required company_id
                cid = record.get("company_id")
                if not self.validator.check_required_company_id(table_name, cid):
                    rows_rejected += 1
                    continue

                # DQ-03 Foreign-key validity: must exist in companies
                if not self.validator.check_fk_validity(
                    table_name, cid, self.valid_company_ids
                ):
                    rows_rejected += 1
                    continue

                # DQ-05 Valid Year / Date
                if "year" in record:
                    self.validator.check_valid_year(
                        table_name, cid, record["year"]
                    )
                if "year" in record and table_name in (
                    "profitandloss",
                    "balancesheet",
                    "cashflow",
                    "financial_ratios",
                ):
                    self.validator.check_duplicate_company_year(
                        table_name, cid, record["year"], seen_pairs
                    )

                # Domain-specific DQ rules
                if table_name == "balancesheet":
                    self.validator.check_critical_values(
                        table_name, cid, record.get("year"), record
                    )
                    self.validator.check_balance_sheet_consistency(
                        cid, record.get("year"), record
                    )
                elif table_name == "profitandloss":
                    self.validator.check_sales_validity(
                        cid, record.get("year"), record.get("sales")
                    )
                    self.validator.check_operating_profit_consistency(
                        cid,
                        record.get("year"),
                        record.get("sales"),
                        record.get("expenses"),
                        record.get("operating_profit"),
                        record.get("opm_percentage"),
                    )
                    if cid and record.get("year") and record.get("net_profit"):
                        pnl_records[(cid, record["year"])] = (
                            record.get("net_profit"),
                            record.get("sales"),
                        )
                elif table_name == "cashflow":
                    self.validator.check_cash_flow_consistency(
                        cid,
                        record.get("year"),
                        record.get("operating_activity"),
                        record.get("investing_activity"),
                        record.get("financing_activity"),
                        record.get("net_cash_flow"),
                    )
                elif table_name == "stock_prices":
                    self.validator.check_critical_values(
                        table_name, cid, record.get("date"), record
                    )
                    self.validator.check_stock_price_logic(
                        cid, record.get("date"), record
                    )

            # Record is referentially clean and ready for DB insertion
            row_tuple = tuple(record[col] for col in target_cols)
            records_to_insert.append(row_tuple)

        # Transactional insertion
        try:
            with conn:
                conn.executemany(insert_sql, records_to_insert)
            rows_inserted = len(records_to_insert)
            status = (
                "SUCCESS"
                if rows_rejected == 0
                else "COMPLETED_WITH_REJECTIONS"
            )
        except Exception as e:
            conn.rollback()
            status = f"FAILED: {str(e)}"
            rows_inserted = 0
            rows_rejected = rows_read

        audit_entry = {
            "source_file": src_filename,
            "source_sheet": sheet_name,
            "target_table": table_name,
            "rows_read": rows_read,
            "rows_inserted": rows_inserted,
            "rows_rejected": rows_rejected,
            "status": status,
            "timestamp": datetime.now().isoformat(),
        }
        self.audit_log.append(audit_entry)
        return audit_entry

    def perform_coverage_and_cross_checks(self, conn: sqlite3.Connection):
        """Run DQ-10 and DQ-15 cross checks after tables are loaded."""
        c = conn.cursor()

        # DQ-15: Year coverage per company in P&L, BS, CF
        for table in ["profitandloss", "balancesheet", "cashflow"]:
            c.execute(
                f"SELECT company_id, GROUP_CONCAT(year) FROM {table} GROUP BY company_id"
            )
            for cid, years_str in c.fetchall():
                years_list = years_str.split(",") if years_str else []
                self.validator.check_year_coverage(table, cid, years_list, min_years=5)

        # DQ-10: Cross statement consistency (P&L net profit margin vs financial_ratios)
        c.execute("""
            SELECT p.company_id, p.year, p.net_profit, p.sales, r.net_profit_margin_pct
            FROM profitandloss p
            JOIN financial_ratios r ON p.company_id = r.company_id AND p.year = r.year
            WHERE p.sales > 0 AND p.net_profit IS NOT NULL AND r.net_profit_margin_pct IS NOT NULL
        """)
        for cid, yr, np, sales, npm in c.fetchall():
            self.validator.check_financial_cross_consistency(cid, yr, np, npm, sales)

    def write_audit_csv(self, file_path: str):
        """Export load_audit.csv."""
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(
                f,
                fieldnames=[
                    "source_file",
                    "source_sheet",
                    "target_table",
                    "rows_read",
                    "rows_inserted",
                    "rows_rejected",
                    "status",
                    "timestamp",
                ],
            )
            writer.writeheader()
            for entry in self.audit_log:
                writer.writerow(entry)

    def generate_manual_review_report(
        self,
        conn: sqlite3.Connection,
        sample_companies: List[str] = ["ABB", "HDFCBANK", "INFY", "RELIANCE", "TCS"],
    ) -> str:
        """Generate a structured markdown/text report for representative companies."""
        c = conn.cursor()
        report_lines = [
            "=" * 78,
            "MANUAL REVIEW REPORT: 5 REPRESENTATIVE COMPANIES COVERAGE",
            "=" * 78,
            f"{'Company':<12} | {'Years (P&L)':<12} | {'P&L Rows':<9} | {'BS Rows':<8} | {'CF Rows':<8} | {'Prices':<8} | {'Ratios':<8}",
            "-" * 78,
        ]

        for cid in sample_companies:
            c.execute("SELECT COUNT(*) FROM profitandloss WHERE company_id = ?", (cid,))
            pnl_cnt = c.fetchone()[0]
            c.execute("SELECT COUNT(DISTINCT year) FROM profitandloss WHERE company_id = ? AND year != 'TTM'", (cid,))
            pnl_yrs = c.fetchone()[0]
            c.execute("SELECT COUNT(*) FROM balancesheet WHERE company_id = ?", (cid,))
            bs_cnt = c.fetchone()[0]
            c.execute("SELECT COUNT(*) FROM cashflow WHERE company_id = ?", (cid,))
            cf_cnt = c.fetchone()[0]
            c.execute("SELECT COUNT(*) FROM stock_prices WHERE company_id = ?", (cid,))
            sp_cnt = c.fetchone()[0]
            c.execute("SELECT COUNT(*) FROM financial_ratios WHERE company_id = ?", (cid,))
            fr_cnt = c.fetchone()[0]

            report_lines.append(
                f"{cid:<12} | {pnl_yrs:<12} | {pnl_cnt:<9} | {bs_cnt:<8} | {cf_cnt:<8} | {sp_cnt:<8} | {fr_cnt:<8}"
            )

        report_lines.append("=" * 78)
        report_str = "\n".join(report_lines)
        return report_str

    def run_pipeline(self) -> Dict[str, Any]:
        """Execute complete ETL pipeline."""
        print(f"Connecting to database: {self.db_path}")
        os.makedirs(os.path.dirname(os.path.abspath(self.db_path)) or ".", exist_ok=True)
        conn = sqlite3.connect(self.db_path)
        conn.execute("PRAGMA foreign_keys = ON;")

        print(f"Applying schema from: {self.schema_path}")
        self.initialize_database(conn)

        print(f"Loading tables from source datasets in: {self.data_dir}")
        for config in TABLE_CONFIGS:
            tname = config["table_name"]
            audit = self.load_table(conn, config)
            print(
                f"  [{tname:<16}] Read: {audit['rows_read']:<5} | Inserted: {audit['rows_inserted']:<5} | Rejected: {audit['rows_rejected']:<4} | Status: {audit['status']}"
            )

        print("Executing cross-table consistency & year coverage validations...")
        self.perform_coverage_and_cross_checks(conn)

        # Integrity Checks
        c = conn.cursor()
        c.execute("PRAGMA foreign_key_check;")
        fk_violations = c.fetchall()
        c.execute("SELECT COUNT(*) FROM companies;")
        companies_count = c.fetchone()[0]

        # Audit and Failure Reports
        audit_csv_path = os.path.join(self.output_dir, "load_audit.csv")
        failures_csv_path = os.path.join(self.output_dir, "validation_failures.csv")

        self.write_audit_csv(audit_csv_path)
        self.validator.export_failures_csv(failures_csv_path)

        # Manual review report
        review_report = self.generate_manual_review_report(conn)
        print("\n" + review_report + "\n")

        summary = self.validator.get_summary()
        print("Data Quality Summary:")
        print(f"  CRITICAL : {summary.get('CRITICAL', 0)}")
        print(f"  WARNING  : {summary.get('WARNING', 0)}")
        print(f"  INFO     : {summary.get('INFO', 0)}")
        print(f"  TOTAL    : {summary.get('TOTAL', 0)}")
        print(f"Companies Count: {companies_count}")
        print(f"Foreign Key Violations: {len(fk_violations)}")

        conn.close()
        return {
            "companies_count": companies_count,
            "fk_violations": fk_violations,
            "summary": summary,
            "audit": self.audit_log,
            "review_report": review_report,
        }


def main():
    parser = argparse.ArgumentParser(description="ETL Pipeline for svgBA100 Financial Intelligence Platform")
    parser.add_argument("--db", default="nifty100.db", help="Path to SQLite database")
    parser.add_argument(
        "--data-dir",
        default="./data",
        help="Path to source Excel files",
    )
    parser.add_argument("--output-dir", default="output", help="Path to output directory")
    parser.add_argument("--schema", default="db/schema.sql", help="Path to schema.sql")

    args = parser.parse_args()
    loader = ETLLoader(
        db_path=args.db,
        data_dir=args.data_dir,
        output_dir=args.output_dir,
        schema_path=args.schema,
    )
    loader.run_pipeline()


if __name__ == "__main__":
    main()




