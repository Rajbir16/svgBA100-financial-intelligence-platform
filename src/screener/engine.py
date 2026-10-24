"""Sprint 3 Financial Screener and Composite Scoring Engine."""

import os
import sqlite3
from typing import Any, Dict, List, Optional, Tuple
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import yaml

DEFAULT_CONFIG_PATH = os.path.join("config", "screener_config.yaml")

KPI_COLUMNS_INFO = [
    ("company_id", "Company ID", "text"),
    ("company_name", "Company Name", "text"),
    ("broad_sector", "Sector", "text"),
    ("sub_sector", "Sub Sector", "text"),
    ("year", "FY Year", "text"),
    ("composite_quality_score", "Quality Score (0-100)", "float"),
    ("return_on_equity_pct", "ROE (%)", "float"),
    ("return_on_capital_employed_pct", "ROCE (%)", "float"),
    ("net_profit_margin_pct", "Net Profit Margin (%)", "float"),
    ("operating_profit_margin_pct", "OPM (%)", "float"),
    ("debt_to_equity", "Debt-to-Equity", "float"),
    ("interest_coverage", "Interest Coverage", "float"),
    ("free_cash_flow_cr", "Free Cash Flow (Cr)", "float"),
    ("sales_cr", "Sales (Cr)", "float"),
    ("net_profit_cr", "Net Profit (Cr)", "float"),
    ("pe_ratio", "P/E Ratio", "float"),
    ("pb_ratio", "P/B Ratio", "float"),
    ("dividend_yield_pct", "Dividend Yield (%)", "float"),
    ("revenue_cagr_5yr", "Rev CAGR 5yr (%)", "float"),
    ("pat_cagr_5yr", "PAT CAGR 5yr (%)", "float"),
]


def load_config(config_path: str = DEFAULT_CONFIG_PATH) -> Dict[str, Any]:
    """Load YAML screener configuration."""
    if not os.path.exists(config_path):
        raise FileNotFoundError(f"Configuration file not found: {config_path}")
    with open(config_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def _load_raw_screener_data(conn: sqlite3.Connection) -> pd.DataFrame:
    """Load latest annual financial data, previous year D/E, and market cap for all 92 companies."""
    query = """
        WITH latest_fr AS (
            SELECT fr.*,
                   ROW_NUMBER() OVER (
                       PARTITION BY fr.company_id
                       ORDER BY 
                           CASE 
                               WHEN fr.company_id = 'SIEMENS' AND fr.year = 'Sep 2024' THEN 1
                               WHEN fr.company_id != 'SIEMENS' AND fr.year = 'Mar 2024' THEN 1
                               WHEN fr.year LIKE '%2024%' AND fr.year != 'TTM' THEN 2
                               ELSE 3
                           END,
                           fr.id DESC
                   ) as rn
            FROM financial_ratios fr
            WHERE fr.sales_cr IS NOT NULL AND fr.year != 'TTM'
        ),
        prev_fr AS (
            SELECT fr.company_id, fr.debt_to_equity AS prev_debt_to_equity,
                   ROW_NUMBER() OVER (
                       PARTITION BY fr.company_id
                       ORDER BY 
                           CASE 
                               WHEN fr.company_id = 'SIEMENS' AND fr.year = 'Sep 2023' THEN 1
                               WHEN fr.company_id != 'SIEMENS' AND fr.year = 'Mar 2023' THEN 1
                               ELSE 3
                           END,
                           fr.id DESC
                   ) as rn
            FROM financial_ratios fr
            WHERE fr.sales_cr IS NOT NULL AND fr.year != 'TTM'
        ),
        fcf_2019 AS (
            SELECT fr.company_id, fr.free_cash_flow_cr AS fcf_start,
                   ROW_NUMBER() OVER (
                       PARTITION BY fr.company_id
                       ORDER BY 
                           CASE 
                               WHEN fr.company_id = 'SIEMENS' AND fr.year = 'Sep 2019' THEN 1
                               WHEN fr.company_id != 'SIEMENS' AND fr.year = 'Mar 2019' THEN 1
                               ELSE 3
                           END,
                           fr.id DESC
                   ) as rn
            FROM financial_ratios fr
            WHERE fr.free_cash_flow_cr IS NOT NULL AND fr.year != 'TTM'
        )
        SELECT 
            c.id AS company_id,
            c.company_name,
            s.broad_sector,
            s.sub_sector,
            l.year,
            l.return_on_equity_pct,
            l.return_on_capital_employed_pct,
            l.net_profit_margin_pct,
            l.operating_profit_margin_pct,
            l.debt_to_equity,
            l.interest_coverage,
            l.icr_label,
            l.asset_turnover,
            l.free_cash_flow_cr,
            l.sales_cr,
            l.net_profit_cr,
            l.dividend_payout_ratio_pct,
            l.revenue_cagr_3yr,
            l.revenue_cagr_5yr,
            l.pat_cagr_5yr,
            l.eps_cagr_5yr,
            l.cfo_quality_5yr,
            mc.pe_ratio,
            mc.pb_ratio,
            mc.dividend_yield_pct,
            mc.market_cap_crore,
            p.prev_debt_to_equity,
            f19.fcf_start
        FROM companies c
        LEFT JOIN sectors s ON c.id = s.company_id
        LEFT JOIN latest_fr l ON c.id = l.company_id AND l.rn = 1
        LEFT JOIN prev_fr p ON c.id = p.company_id AND p.rn = 1
        LEFT JOIN fcf_2019 f19 ON c.id = f19.company_id AND f19.rn = 1
        LEFT JOIN market_cap mc ON c.id = mc.company_id AND mc.year = 2024
    """
    return pd.read_sql_query(query, conn)


def _scale_series_sector_relative(df: pd.DataFrame, metric_col: str, group_col: str = "broad_sector", invert: bool = False) -> pd.Series:
    """Scale a metric between 0 and 100 sector-relative using P10/P90 winsorisation."""
    scaled_series = pd.Series(50.0, index=df.index, dtype=float)
    
    # Universe fallback percentiles
    univ_vals = df[metric_col].dropna()
    if len(univ_vals) > 0:
        univ_p10 = float(np.percentile(univ_vals, 10))
        univ_p90 = float(np.percentile(univ_vals, 90))
    else:
        univ_p10, univ_p90 = 0.0, 100.0

    for sector, group in df.groupby(group_col):
        valid = group[metric_col].dropna()
        if len(valid) >= 3:
            p10 = float(np.percentile(valid, 10))
            p90 = float(np.percentile(valid, 90))
        else:
            p10, p90 = univ_p10, univ_p90

        for idx in group.index:
            val = df.loc[idx, metric_col]
            if pd.isna(val) or val is None:
                scaled_series.loc[idx] = 50.0
            else:
                val_float = float(val)
                clipped = min(max(val_float, p10), p90)
                if p90 > p10:
                    score = (clipped - p10) / (p90 - p10) * 100.0
                else:
                    score = 50.0
                
                if invert:
                    score = 100.0 - score
                scaled_series.loc[idx] = max(0.0, min(100.0, score))

    return scaled_series


def compute_composite_quality_scores(df: pd.DataFrame) -> pd.Series:
    """Compute 0-100 Composite Quality Score with sector-relative P10/P90 normalization.
    
    Weights:
      Profitability (35%): ROE (15), ROCE (10), NPM (10)
      Cash Quality (30%): FCF CAGR (15), CFO/PAT (10), FCF positive (5)
      Growth (20%): Revenue CAGR 5yr (10), PAT CAGR 5yr (10)
      Leverage (15%): D/E score (10), ICR score (5)
    """
    # 1. Profitability components
    s_roe = _scale_series_sector_relative(df, "return_on_equity_pct")
    s_roce = _scale_series_sector_relative(df, "return_on_capital_employed_pct")
    s_npm = _scale_series_sector_relative(df, "net_profit_margin_pct")

    # 2. Cash Quality components
    # Compute 5-year FCF CAGR
    fcf_cagr_list = []
    for _, row in df.iterrows():
        f_end = row.get("free_cash_flow_cr")
        f_start = row.get("fcf_start")
        if pd.notna(f_end) and pd.notna(f_start):
            if f_start > 0 and f_end > 0:
                cagr = ((float(f_end) / float(f_start)) ** 0.2 - 1.0) * 100.0
            elif f_end > 0 and f_start <= 0:
                cagr = 25.0  # Turnaround bonus
            elif f_end <= 0 and f_start > 0:
                cagr = -25.0  # Deterioration penalty
            else:
                cagr = -10.0
        elif pd.notna(f_end) and f_end > 0:
            cagr = 10.0
        else:
            cagr = 0.0
        fcf_cagr_list.append(cagr)
    
    df_temp = df.copy()
    df_temp["fcf_cagr_5yr"] = fcf_cagr_list
    s_fcf_cagr = _scale_series_sector_relative(df_temp, "fcf_cagr_5yr")
    s_cfo_pat = _scale_series_sector_relative(df, "cfo_quality_5yr")
    s_fcf_pos = df["free_cash_flow_cr"].apply(lambda x: 100.0 if (pd.notna(x) and x > 0) else 0.0)

    # 3. Growth components
    s_rev_cagr = _scale_series_sector_relative(df, "revenue_cagr_5yr")
    s_pat_cagr = _scale_series_sector_relative(df, "pat_cagr_5yr")

    # 4. Leverage components
    # D/E: lower is better (inverted). Financials are exempted / given benchmark high score.
    s_de_raw = _scale_series_sector_relative(df, "debt_to_equity", invert=True)
    s_de = []
    for idx, row in df.iterrows():
        if str(row.get("broad_sector")).lower() == "financials":
            s_de.append(85.0)  # Standard high score for financial institutions
        else:
            s_de.append(s_de_raw.loc[idx])
    s_de = pd.Series(s_de, index=df.index)

    # ICR: Debt free is treated as infinite coverage -> 100.0
    s_icr_raw = _scale_series_sector_relative(df, "interest_coverage")
    s_icr = []
    for idx, row in df.iterrows():
        icr_lbl = str(row.get("icr_label") or "")
        icr_val = row.get("interest_coverage")
        if icr_lbl == "Debt Free" or pd.isna(icr_val) or icr_val is None:
            s_icr.append(100.0)
        else:
            s_icr.append(s_icr_raw.loc[idx])
    s_icr = pd.Series(s_icr, index=df.index)

    # Combine with weights (sum = 100)
    composite = (
        0.15 * s_roe +
        0.10 * s_roce +
        0.10 * s_npm +
        0.15 * s_fcf_cagr +
        0.10 * s_cfo_pat +
        0.05 * s_fcf_pos +
        0.10 * s_rev_cagr +
        0.10 * s_pat_cagr +
        0.10 * s_de +
        0.05 * s_icr
    )
    # Clamp strictly to [0.0, 100.0] and round to 2 decimal places
    return composite.clip(0.0, 100.0).round(2)


def get_screener_universe(db_path: str = "nifty100.db") -> pd.DataFrame:
    """Retrieve full screener dataset with computed Composite Quality Score."""
    conn = sqlite3.connect(db_path)
    df = _load_raw_screener_data(conn)
    
    # Calculate composite quality scores
    df["composite_quality_score"] = compute_composite_quality_scores(df)
    
    # Update financial_ratios table in DB with the 0-100 composite score for latest rows
    for _, r in df.iterrows():
        conn.execute("""
            UPDATE financial_ratios 
            SET composite_quality_score = ?
            WHERE company_id = ? AND year = ?
        """, (float(r["composite_quality_score"]), r["company_id"], r["year"]))
    conn.commit()
    conn.close()
    return df


def apply_screener_filters(df: pd.DataFrame, filters: Dict[str, Any]) -> pd.DataFrame:
    """Filter universe according to the 15 metrics, preset rules, and edge cases."""
    res = df.copy()
    
    for key, val in filters.items():
        if val is None or key in ("description", "name"):
            continue

        if key == "roe_min":
            res = res[res["return_on_equity_pct"].notna() & (res["return_on_equity_pct"] > float(val))]
        elif key == "debt_to_equity_max":
            # Rule: D/E filter skips Financials sector!
            is_fin = res["broad_sector"].str.lower() == "financials"
            de_ok = res["debt_to_equity"].notna() & (res["debt_to_equity"] <= float(val))
            res = res[is_fin | de_ok]
        elif key == "debt_to_equity_exact":
            is_fin = res["broad_sector"].str.lower() == "financials"
            de_ok = res["debt_to_equity"].fillna(0.0) == float(val)
            res = res[is_fin | de_ok]
        elif key == "fcf_min":
            res = res[res["free_cash_flow_cr"].notna() & (res["free_cash_flow_cr"] > float(val))]
        elif key == "revenue_cagr_5yr_min":
            res = res[res["revenue_cagr_5yr"].notna() & (res["revenue_cagr_5yr"] > float(val))]
        elif key == "revenue_cagr_3yr_min":
            res = res[res["revenue_cagr_3yr"].notna() & (res["revenue_cagr_3yr"] > float(val))]
        elif key == "pat_cagr_5yr_min":
            res = res[res["pat_cagr_5yr"].notna() & (res["pat_cagr_5yr"] > float(val))]
        elif key == "opm_min":
            res = res[res["operating_profit_margin_pct"].notna() & (res["operating_profit_margin_pct"] > float(val))]
        elif key == "pe_max":
            res = res[res["pe_ratio"].notna() & (res["pe_ratio"] <= float(val))]
        elif key == "pb_max":
            res = res[res["pb_ratio"].notna() & (res["pb_ratio"] <= float(val))]
        elif key == "dividend_yield_min":
            res = res[res["dividend_yield_pct"].notna() & (res["dividend_yield_pct"] >= float(val))]
        elif key == "dividend_payout_max":
            res = res[res["dividend_payout_ratio_pct"].notna() & (res["dividend_payout_ratio_pct"] < float(val))]
        elif key == "icr_min":
            # Rule: Debt Free ICR = infinity
            is_debt_free = (res["icr_label"] == "Debt Free") | res["interest_coverage"].isna()
            icr_ok = res["interest_coverage"].notna() & (res["interest_coverage"] >= float(val))
            res = res[is_debt_free | icr_ok]
        elif key == "market_cap_min":
            res = res[res["market_cap_crore"].notna() & (res["market_cap_crore"] >= float(val))]
        elif key == "net_profit_min":
            res = res[res["net_profit_cr"].notna() & (res["net_profit_cr"] >= float(val))]
        elif key == "eps_cagr_min":
            res = res[res["eps_cagr_5yr"].notna() & (res["eps_cagr_5yr"] >= float(val))]
        elif key == "asset_turnover_min":
            res = res[res["asset_turnover"].notna() & (res["asset_turnover"] >= float(val))]
        elif key == "sales_min":
            res = res[res["sales_cr"].notna() & (res["sales_cr"] > float(val))]
        elif key == "de_declining_yoy" and val:
            # D/E declining YoY: de < prev_debt_to_equity OR de == 0 OR Financials
            is_fin = res["broad_sector"].str.lower() == "financials"
            de_dec = (res["debt_to_equity"].notna() & res["prev_debt_to_equity"].notna() & (res["debt_to_equity"] < res["prev_debt_to_equity"])) | (res["debt_to_equity"] == 0.0)
            res = res[is_fin | de_dec]

    # Return results sorted by composite_quality_score descending
    return res.sort_values(by="composite_quality_score", ascending=False)


def run_presets(df: pd.DataFrame, config: Dict[str, Any]) -> Dict[str, pd.DataFrame]:
    """Execute all 6 screener presets and return results dictionary."""
    presets = config.get("presets", {})
    results = {}
    for preset_name, preset_filters in presets.items():
        results[preset_name] = apply_screener_filters(df, preset_filters)
    return results


def generate_screener_excel(df: pd.DataFrame, config: Dict[str, Any], output_path: str = "output/screener_output.xlsx") -> None:
    """Generate screener_output.xlsx with exactly 6 preset sheets and formatted 20 KPI columns."""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    preset_results = run_presets(df, config)
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    green_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    green_font = Font(name="Calibri", size=10, color="155724")
    
    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    
    for preset_name, p_df in preset_results.items():
        ws = wb.create_sheet(title=preset_name[:31])
        ws.views.sheetView[0].showGridLines = True
        
        headers = [col_info[1] for col_info in KPI_COLUMNS_INFO]
        ws.append(headers)
        ws.row_dimensions[1].height = 28
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = navy_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            
        filters = config.get("presets", {}).get(preset_name, {})
        
        for r_idx, (_, row) in enumerate(p_df.iterrows(), start=2):
            row_vals = []
            for col_field, _, col_type in KPI_COLUMNS_INFO:
                val = row.get(col_field)
                if pd.isna(val) or val is None:
                    row_vals.append("N/A")
                elif col_type == "float":
                    row_vals.append(round(float(val), 2))
                else:
                    row_vals.append(str(val))
                    
            ws.append(row_vals)
            ws.row_dimensions[r_idx].height = 20
            
            for c_idx in range(1, len(row_vals) + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                cell.font = data_font
                cell.alignment = Alignment(horizontal="center" if c_idx not in (2, 3, 4) else "left", vertical="center")
                
                # Check if this cell is one of the meeting criteria to shade green
                field_name = KPI_COLUMNS_INFO[c_idx - 1][0]
                cell_val = row.get(field_name)
                
                met_criteria = False
                if field_name == "return_on_equity_pct" and "roe_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) > filters["roe_min"])
                elif field_name == "debt_to_equity" and "debt_to_equity_max" in filters:
                    is_fin = str(row.get("broad_sector")).lower() == "financials"
                    met_criteria = is_fin or (pd.notna(cell_val) and float(cell_val) <= filters["debt_to_equity_max"])
                elif field_name == "free_cash_flow_cr" and "fcf_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) > filters["fcf_min"])
                elif field_name == "revenue_cagr_5yr" and "revenue_cagr_5yr_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) > filters["revenue_cagr_5yr_min"])
                elif field_name == "pat_cagr_5yr" and "pat_cagr_5yr_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) > filters["pat_cagr_5yr_min"])
                elif field_name == "pe_ratio" and "pe_max" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) <= filters["pe_max"])
                elif field_name == "pb_ratio" and "pb_max" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) <= filters["pb_max"])
                elif field_name == "dividend_yield_pct" and "dividend_yield_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) >= filters["dividend_yield_min"])
                elif field_name == "sales_cr" and "sales_min" in filters and pd.notna(cell_val):
                    met_criteria = (float(cell_val) > filters["sales_min"])
                elif field_name == "composite_quality_score":
                    cell.font = bold_font
                    
                if met_criteria:
                    cell.fill = green_fill
                    cell.font = green_font

        # Auto adjust column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)


def main() -> None:
    config = load_config()
    print("Loading universe data & computing composite quality score...")
    df = get_screener_universe()
    print(f"Loaded {len(df)} companies in screener universe.")
    
    presets = run_presets(df, config)
    print("\n--- Screener Preset Results ---")
    for name, res_df in presets.items():
        print(f"{name}: {len(res_df)} companies")
        
    print("\nGenerating screener_output.xlsx...")
    generate_screener_excel(df, config)
    print("screener_output.xlsx generated successfully.")


if __name__ == "__main__":
    main()

