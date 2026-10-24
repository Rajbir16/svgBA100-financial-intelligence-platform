"""Fourteen Sprint 3 DQ rule tests for Screener, Peer Engine, Scoring, and Reporting."""

import os
import sqlite3
import openpyxl
import pandas as pd
import pytest

from src.analytics.peer import (
    calculate_percentiles_for_group,
    compute_peer_percentiles,
    get_company_peer_percentiles,
)
from src.screener.engine import (
    apply_screener_filters,
    get_screener_universe,
    load_config,
    run_presets,
)


@pytest.fixture(scope="module")
def screener_data():
    config = load_config()
    df = get_screener_universe("nifty100.db")
    presets = run_presets(df, config)
    return config, df, presets


def test_dq_s01_config_structure(screener_data):
    """DQ-S1: Configuration structure and 15 filters validation."""
    config, _, _ = screener_data
    assert "filters" in config
    assert "presets" in config
    assert "composite_score" in config
    assert len(config["filters"]) >= 15
    assert len(config["presets"]) == 6
    preset_names = list(config["presets"].keys())
    expected = [
        "Quality Compounder",
        "Value Pick",
        "Growth Accelerator",
        "Dividend Champion",
        "Debt-Free Blue Chip",
        "Turnaround Watch",
    ]
    for exp in expected:
        assert exp in preset_names


def test_dq_s02_preset_quality_compounder(screener_data):
    """DQ-S2: Quality Compounder returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Quality Compounder"]
    count = len(res)
    assert 5 <= count <= 50, f"Quality Compounder count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["return_on_equity_pct"] > 15.0
        assert row["free_cash_flow_cr"] > 0.0
        assert row["revenue_cagr_5yr"] > 10.0
        if str(row["broad_sector"]).lower() != "financials":
            assert row["debt_to_equity"] < 1.0


def test_dq_s03_preset_value_pick(screener_data):
    """DQ-S3: Value Pick returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Value Pick"]
    count = len(res)
    assert 5 <= count <= 50, f"Value Pick count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["pe_ratio"] <= 20.0
        assert row["pb_ratio"] <= 4.5
        assert row["dividend_yield_pct"] >= 1.0
        if str(row["broad_sector"]).lower() != "financials":
            assert row["debt_to_equity"] <= 2.0


def test_dq_s04_preset_growth_accelerator(screener_data):
    """DQ-S4: Growth Accelerator returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Growth Accelerator"]
    count = len(res)
    assert 5 <= count <= 50, f"Growth Accelerator count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["pat_cagr_5yr"] > 20.0
        assert row["revenue_cagr_5yr"] > 15.0
        if str(row["broad_sector"]).lower() != "financials":
            assert row["debt_to_equity"] <= 2.0


def test_dq_s05_preset_dividend_champion(screener_data):
    """DQ-S5: Dividend Champion returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Dividend Champion"]
    count = len(res)
    assert 5 <= count <= 50, f"Dividend Champion count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["dividend_yield_pct"] >= 2.0
        assert row["dividend_payout_ratio_pct"] < 80.0
        assert row["free_cash_flow_cr"] > 0.0


def test_dq_s06_preset_debt_free_blue_chip(screener_data):
    """DQ-S6: Debt-Free Blue Chip returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Debt-Free Blue Chip"]
    count = len(res)
    assert 5 <= count <= 50, f"Debt-Free Blue Chip count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["return_on_equity_pct"] > 12.0
        assert row["sales_cr"] > 5000.0
        if str(row["broad_sector"]).lower() != "financials":
            assert row["debt_to_equity"] <= 0.05


def test_dq_s07_preset_turnaround_watch(screener_data):
    """DQ-S7: Turnaround Watch returns 5-50 companies and matches criteria."""
    _, _, presets = screener_data
    res = presets["Turnaround Watch"]
    count = len(res)
    assert 5 <= count <= 50, f"Turnaround Watch count {count} not in [5, 50]"
    for _, row in res.iterrows():
        assert row["revenue_cagr_3yr"] > 10.0
        assert row["free_cash_flow_cr"] > 0.0


def test_dq_s08_financials_de_filter_rule(screener_data):
    """DQ-S8: Financials sector companies bypass D/E filters."""
    _, df, _ = screener_data
    fin_df = df[df["broad_sector"].str.lower() == "financials"]
    assert len(fin_df) > 0
    # Apply strict D/E filter that would fail almost all banks if applied
    filtered = apply_screener_filters(fin_df, {"debt_to_equity_max": 0.5})
    # Since rule skips Financials, all financials remain
    assert len(filtered) == len(fin_df)


def test_dq_s09_debt_free_icr_infinite_rule(screener_data):
    """DQ-S9: Debt free companies with None interest coverage pass ICR threshold."""
    _, df, _ = screener_data
    # Create test sample with a debt free company having None interest coverage
    sample_df = df.head(5).copy()
    sample_df.loc[sample_df.index[0], "interest_coverage"] = None
    sample_df.loc[sample_df.index[0], "icr_label"] = "Debt Free"
    debt_free_id = sample_df.loc[sample_df.index[0], "company_id"]

    # Apply strict ICR requirement of 50.0
    filtered = apply_screener_filters(sample_df, {"icr_min": 50.0})
    assert debt_free_id in filtered["company_id"].values



def test_dq_s10_composite_score_bounds_and_sector_relative(screener_data):
    """DQ-S10: Composite Quality Score is in [0, 100] and non-null for all 92 companies."""
    _, df, _ = screener_data
    assert len(df) == 92
    scores = df["composite_quality_score"]
    assert scores.notna().all()
    assert (scores >= 0.0).all()
    assert (scores <= 100.0).all()
    assert scores.max() > scores.min()


def test_dq_s11_peer_percentiles_bounds_and_inversion():
    """DQ-S11: Peer percentiles are in [0, 100] and D/E is properly inverted."""
    conn = sqlite3.connect("nifty100.db")
    df_pct = pd.read_sql_query("SELECT * FROM peer_percentiles", conn)
    conn.close()
    assert len(df_pct) == 560  # 56 companies * 10 metrics
    assert (df_pct["percentile_rank"] >= 0.0).all()
    assert (df_pct["percentile_rank"] <= 100.0).all()

    # Synthetic test to verify inversion logic
    mock_group = pd.DataFrame({
        "company_id": ["A", "B", "C"],
        "debt_to_equity": [0.1, 0.5, 1.2],
    })
    res = calculate_percentiles_for_group(mock_group, "debt_to_equity", invert=True)
    res_map = {r[0]: r[2] for r in res}
    # Company A has lowest D/E -> must have highest percentile rank
    assert res_map["A"] > res_map["B"] > res_map["C"]


def test_dq_s12_peer_rankings_it_services_and_fmcg():
    """DQ-S12: IT Services highest ROE has highest percentile rank; FMCG ranking verification."""
    conn = sqlite3.connect("nifty100.db")
    it_df = pd.read_sql_query("""
        SELECT company_id, percentile_rank 
        FROM peer_percentiles 
        WHERE peer_group_name = 'IT Services' AND metric = 'ROE'
        ORDER BY percentile_rank DESC
    """, conn)
    fmcg_df = pd.read_sql_query("""
        SELECT company_id, percentile_rank 
        FROM peer_percentiles 
        WHERE peer_group_name = 'FMCG' AND metric = 'ROE'
        ORDER BY percentile_rank DESC
    """, conn)
    conn.close()

    # In IT Services, TCS has highest ROE and must be top ranked (100.0%)
    assert it_df.iloc[0]["company_id"] == "TCS"
    assert it_df.iloc[0]["percentile_rank"] == 100.0

    # In FMCG, NESTLEIND has highest ROE
    assert fmcg_df.iloc[0]["company_id"] == "NESTLEIND"
    assert fmcg_df.iloc[0]["percentile_rank"] == 100.0


def test_dq_s13_unassigned_peer_group_handling():
    """DQ-S13: Companies without peer groups return 'No peer group assigned' without crashing."""
    res1 = get_company_peer_percentiles("RECLTD")
    assert res1 == "No peer group assigned"

    res2 = get_company_peer_percentiles("SIEMENS")
    assert res2 == "No peer group assigned"

    # Known peer group company returns dict
    res_tcs = get_company_peer_percentiles("TCS")
    assert isinstance(res_tcs, dict)
    assert res_tcs["peer_group_name"] == "IT Services"


def test_dq_s14_output_artifacts_validation():
    """DQ-S14: Verify screener Excel has 6 sheets, peer Excel has 11 sheets, and 92 radar charts exist."""
    assert os.path.exists("output/screener_output.xlsx")
    wb_screen = openpyxl.load_workbook("output/screener_output.xlsx")
    assert len(wb_screen.sheetnames) == 6

    assert os.path.exists("output/peer_comparison.xlsx")
    wb_peer = openpyxl.load_workbook("output/peer_comparison.xlsx")
    assert len(wb_peer.sheetnames) == 11

    chart_files = [f for f in os.listdir("reports/radar_charts") if f.endswith(".png")]
    assert len(chart_files) == 92
